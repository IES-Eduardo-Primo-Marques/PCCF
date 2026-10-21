#!/usr/bin/env python3
"""Regenera libro_{CICLE}.xlsx des del JSON sense perdre el treball dels docents.

La generació normal (json2excel.py) no sobreescriu mai un llibre existent. Quan
canvien els RA/CE d'algun mòdul al JSON, cal regenerar el llibre i traslladar-hi
el que ja han omplit els docents a la resta de fulles:

1. Genera un llibre nou des del JSON (json2excel.py) en un directori temporal.
2. Per a cada fulla del llibre nou que també existix al vell:
   - si els RA (columna B) i els CE (columna E) coincidixen, copia del vell les
     cel·les editables (les desbloquejades al llibre nou: % RA, HORES, % CE,
     REQUISIT FE, HORES DUAL, CONTINGUTS), casant les columnes pel títol de la
     fila de capçalera, i també les alçades de fila;
   - si no coincidixen, la fulla s'ha de passar a --regenera (es queda com
     l'acabada de generar, buida) o, si té el mateix nombre de RA, a
     --conserva-ra (conserva % RA i CONTINGUTS de cada RA, i les cel·les dels
     CE que tenen el mateix text dins del mateix RA); si no, no s'escriu res.
3. Amb --pd, actualitza també les PD (PD_{CICLE}_{CODI}_*.md); la resta del
   que ha escrit el docent no es toca:
   - mòduls regenerats o de --conserva-ra: la taula de RA i, si el mòdul ha passat a ser un àmbit
     de grau bàsic (competències específiques CE1…), els textos de RA de la
     plantilla i el nom del mòdul;
   - hores del mòdul («És un mòdul de N hores») si han canviat al JSON;
   - mòduls que ja no dualitzen: elimina la secció «Formació en empresa»;
   - estàndards de competència (UnidadesCompetenciaAcreditadas) si han canviat
     al JSON: es torna a generar la secció «Relació entre els estàndards…»;
   - si el mòdul ha canviat de nom, reanomena el fitxer com ho faria json2pccf.

Per defecte només simula. Amb --aplica fa còpia del llibre i de les PD a
--backup-dir i substituïx els originals.

Ús (al contenidor, sempre amb el PCCF_ROOT d'OneDrive):
    python3 tools/regenera_llibre.py EI SCO --root /home/PCCF/pccf_sync --regenera PA --pd
    python3 tools/regenera_llibre.py EI SCO --root /home/PCCF/pccf_sync --regenera PA --pd --aplica
"""

import argparse
import datetime
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import warnings

import jinja2
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from excel_estils import FILA_CAPCALERA, FILES_BANDA1  # noqa: E402
from pccf_utils import HOJA_INSTRUCCIONS, dualitza, es_ambit, es_fila_ra, get_hoja_label  # noqa: E402

warnings.filterwarnings("ignore")

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
COL_RA, COL_CE = 2, 5
CE_LLETRA_RE = re.compile(r"^[a-zñ]\) *")
RA_FILA_RE = re.compile(r"^\| *(RA|CE)\d+ *\|")
PD_ESTAT_RE = re.compile(r"_(BORRADOR|OK)\.md$")
# «És un mòdul de {{ modulo.horas }} hores» de la plantilla
HORES_RE = re.compile(r"^(És un mòdul de )(\d+)( hores)", re.M)
# Secció «Formació en empresa» de la plantilla (només als mòduls que dualitzen):
# fins al títol següent, «### Pèrdua de l'avaluació contínua».
FE_SECCIO_RE = re.compile(r"^### Formació en empresa \(RA dualitzats\)\n.*?(?=^### )", re.M | re.S)
# Textos de la plantilla de PD que canvien als àmbits: {% if ambit %}NOU{% else %}VELL{% endif %}.
# Els que porten variables (la frase amb el nom del mòdul) van a INTRO_RA_RE.
_AMBIT_RE = re.compile(r"\{% if ambit %\}(.*?)\{% else %\}(.*?)\{% endif %\}", re.S)
with open(os.path.join(REPO, "templates", "_base_pd.md"), encoding="utf-8") as _f:
    TEXTOS_AMBIT = sorted(((vell, nou) for nou, vell in _AMBIT_RE.findall(_f.read())
                           if "{" not in vell + nou), key=lambda p: -len(p[0]))
# Secció «Relació entre els estàndards de competència…» de la plantilla, fins al
# títol següent. El cos es genera amb el mateix tros de Jinja de _base_pd.md.
UC_TITOL = "## Relació entre els estàndards de competència i els mòduls del cicle formatiu\n"
UC_SECCIO_RE = re.compile(r"^" + re.escape(UC_TITOL) + r".*?(?=^## )", re.M | re.S)
UC_CODI_RE = re.compile(r"^\| *(UC_?\d+_\d) *\|", re.M)
with open(os.path.join(REPO, "templates", "_base_pd.md"), encoding="utf-8") as _f:
    _m = re.search(re.escape(UC_TITOL) + r"(.*?\{% endif %\})\{% endif %\}\n\n", _f.read(), re.S)
    PLANTILLA_UC = jinja2.Template(UC_TITOL + _m.group(1) + "\n\n", keep_trailing_newline=True)
INTRO_RA_RE = re.compile(r"^Els \*\*resultats d'aprenentatge\*\* relatius al mòdul de .* són:$", re.M)


def textos(ws):
    """[(fila, RA, CE)] de les files de dades: el que ve del JSON."""
    files = []
    for r in range(FILA_CAPCALERA + 1, ws.max_row + 1):
        ra = ws.cell(row=r, column=COL_RA).value
        ce = ws.cell(row=r, column=COL_CE).value
        ra = ra.strip() if es_fila_ra(ra) else None
        ce = ce.strip() if isinstance(ce, str) and ce.strip() else None
        if ra or ce:
            files.append((r, ra, ce))
    return files


def capcalera(ws):
    """{títol de columna: columna} de la fila de capçalera."""
    cols = {}
    for c in ws[FILA_CAPCALERA]:
        if isinstance(c.value, str) and c.value.strip():
            cols.setdefault(" ".join(c.value.split()), c.column)
    return cols


def es_editable(cell):
    return cell.protection is not None and cell.protection.locked is False


def migra_fulla(vell, nou, files=None):
    """Copia a `nou` els valors editables de `vell`. Torna el nombre de cel·les
    amb valor copiades i les columnes editables que el vell no té. `files`
    ({fila nova: fila vella}) limita la còpia a eixes files; sense, la mateixa
    fila."""
    cap_v, cap_n = capcalera(vell), capcalera(nou)
    titol_n = {col: t for t, col in cap_n.items()}
    copiades, sense_origen = 0, set()
    for row in nou.iter_rows(min_row=FILA_CAPCALERA + 1, max_row=nou.max_row):
        if files is not None and row[0].row not in files:
            continue
        fila_v = files[row[0].row] if files is not None else row[0].row
        for cell in row:
            if not es_editable(cell) or type(cell).__name__ == "MergedCell":
                continue
            titol = titol_n.get(cell.column)
            col_v = cap_v.get(titol) if titol else cell.column
            if col_v is None:
                sense_origen.add(titol)
                continue
            valor = vell.cell(row=fila_v, column=col_v).value
            if valor is not None:
                cell.value = valor
                copiades += 1
    for r_n in (files or {r: r for r in vell.row_dimensions}):
        dim = vell.row_dimensions.get(files[r_n] if files else r_n)
        if dim is not None and dim.height is not None and r_n > FILA_CAPCALERA:
            nou.row_dimensions[r_n].height = dim.height
    return copiades, sense_origen


def _blocs_ra(ws):
    """[(fila del RA, {text del CE sense la lletra: fila})] de la fulla."""
    blocs = []
    for r, ra, ce in textos(ws):
        if ra:
            blocs.append((r, {}))
        elif ce and blocs:
            blocs[-1][1][CE_LLETRA_RE.sub("", ce)] = r
    return blocs


def files_conserva_ra(vell, nou):
    """{fila nova: fila vella} per a --conserva-ra: cada fila de RA amb la del
    mateix RA (% RA, CONTINGUTS) i cada CE amb el del mateix RA que té el mateix
    text (HORES, % CE, FE...). Els CE nous queden buits. None si el nombre de
    RA no coincidix."""
    blocs_v, blocs_n = _blocs_ra(vell), _blocs_ra(nou)
    if len(blocs_v) != len(blocs_n):
        return None, 0
    files, ce_conservats = {}, 0
    for (r_v, ce_v), (r_n, ce_n) in zip(blocs_v, blocs_n):
        files[r_n] = r_v
        for text, fila in ce_n.items():
            if text in ce_v:
                files[fila] = ce_v[text]
                ce_conservats += 1
    return files, ce_conservats


def genera_llibre(cicle, familia, outdir):
    subprocess.run([sys.executable, os.path.join(REPO, "tools", "json2excel.py"),
                    cicle, familia, "--outdir", outdir],
                   cwd=REPO, check=True, stdout=subprocess.DEVNULL)
    return os.path.join(outdir, f"libro_{cicle}.xlsx")


def taula_ra(modul):
    return [f"| {ra} | {v['Resultado']} |" for ra, v in modul["ResultadosAprendizaje"].items()]


def nom_fitxer_modul(nom):
    """El mateix que json2pccf.py::nom_fitxer_modul()."""
    nom = nom.replace(" ", "")
    for a, b in (("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u")):
        nom = nom.replace(a, b)
    return nom


def nom_vell(ws):
    """Nom del mòdul a la capçalera de la fulla vella (MÒDUL, banda 1)."""
    valor = ws.cell(row=FILES_BANDA1[1], column=2).value
    return valor.strip() if isinstance(valor, str) and valor.strip() else None


def actualitza_pd(text, modul, nom_anterior=None):
    """Substituïx les files de la taula de RA (i, si és un àmbit, els textos de
    RA de la plantilla i el nom del mòdul). Torna el text nou o None si no troba
    la taula."""
    if es_ambit(modul):
        for vell, nou in TEXTOS_AMBIT:
            text = text.replace(vell, nou)
        text = INTRO_RA_RE.sub(
            f"Les **competències específiques** de l'àmbit de {modul['nombre']} (annex V del Reial "
            f"decret 217/2022, d'acord amb l'article 3.3 del Decret 117/2025) són:", text)
        if nom_anterior and nom_anterior != modul["nombre"]:
            # «… I» no ha de tocar «… II»
            text = re.sub(re.escape(nom_anterior) + r"(?![\wI])", modul["nombre"], text)
    linies = text.split("\n")
    idx = [i for i, l in enumerate(linies) if RA_FILA_RE.match(l)]
    if not idx or idx != list(range(idx[0], idx[-1] + 1)):
        return None
    linies[idx[0]:idx[-1] + 1] = taula_ra(modul)
    return "\n".join(linies)


def seccio_uc(modul):
    """Secció d'estàndards de competència tal com la genera la plantilla."""
    return PLANTILLA_UC.render(modulo=modul)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("cicle", help="Codi del cicle (EI, IS, FPBIIO...)")
    ap.add_argument("familia", help="Família (INF, SCO)")
    ap.add_argument("--root", default=".", help="PCCF_ROOT (al contenidor: /home/PCCF/pccf_sync)")
    ap.add_argument("--regenera", nargs="*", default=[],
                    help="Fulles que es regeneren buides (RA/CE canviats), p. ex. PA")
    ap.add_argument("--conserva-ra", nargs="*", default=[],
                    help="Fulles amb RA/CE canviats però el mateix nombre de RA: conserva %% RA i "
                         "CONTINGUTS de cada RA i els CE amb el mateix text; la resta queda buida")
    ap.add_argument("--pd", action="store_true",
                    help="Actualitza també la taula de RA de les PD de les fulles regenerades")
    ap.add_argument("--aplica", action="store_true", help="Escriu els canvis (sense: només simula)")
    ap.add_argument("--backup-dir", default=os.path.join(REPO, "temp", "backups_llibres"))
    args = ap.parse_args()

    cicle, familia = args.cicle.upper(), args.familia.upper()
    pdir = os.path.join(args.root, "programacions", cicle)
    llibre = os.path.join(pdir, f"libro_{cicle}.xlsx")
    if not os.path.isfile(llibre):
        sys.exit(f"No existix {llibre}")
    with open(os.path.join(REPO, f"boe_{familia}", f"rd-{cicle.lower()}.json"), encoding="utf-8") as f:
        moduls = json.load(f)["ModulosProfesionales"]
    per_fulla = {get_hoja_label(m["nombre"]): (codi, m) for codi, m in moduls.items()}

    tmp = tempfile.mkdtemp(prefix=f"regenera_{cicle}_")
    nou_path = genera_llibre(cicle, familia, tmp)
    wb_v = load_workbook(llibre)
    wb_n = load_workbook(nou_path)

    regenera, conserva = set(args.regenera), set(args.conserva_ra)
    desconegudes = (regenera | conserva) - set(wb_n.sheetnames)
    if desconegudes:
        sys.exit(f"Fulles de --regenera/--conserva-ra que no existixen: {', '.join(sorted(desconegudes))}")
    if regenera & conserva:
        sys.exit(f"Fulles a --regenera i a --conserva-ra alhora: {', '.join(sorted(regenera & conserva))}")
    canvien = regenera | conserva  # fulles amb RA/CE nous: la taula de RA de la PD també

    errors = []
    print(f"Llibre: {llibre}")
    for nom in wb_n.sheetnames:
        if nom == HOJA_INSTRUCCIONS:
            continue
        nou = wb_n[nom]
        if nom not in wb_v.sheetnames:
            print(f"  {nom:8} NOVA (no era al llibre vell): es queda buida")
            continue
        vell = wb_v[nom]
        iguals = [(ra, ce) for _, ra, ce in textos(vell)] == [(ra, ce) for _, ra, ce in textos(nou)]
        if nom in regenera:
            print(f"  {nom:8} REGENERADA ({'RA/CE iguals que abans!' if iguals else 'RA/CE nous'}): "
                  f"es perd el que tenia")
            continue
        if nom in conserva:
            files, ce_conservats = files_conserva_ra(vell, nou)
            if files is None:
                errors.append(nom)
                print(f"  {nom:8} ERROR: el nombre de RA ha canviat; no es pot usar --conserva-ra (cal --regenera)")
                continue
            copiades, sense = migra_fulla(vell, nou, files)
            print(f"  {nom:8} RA/CE nous, conservat per RA: {copiades} cel·les copiades; "
                  f"{ce_conservats} CE amb el mateix text conservats, la resta buits")
            continue
        if not iguals:
            errors.append(nom)
            print(f"  {nom:8} ERROR: els RA/CE no coincidixen amb el JSON; afig-la a --regenera "
                  f"si s'ha de buidar")
            continue
        copiades, sense = migra_fulla(vell, nou)
        avis = f" (sense columna al vell: {', '.join(sorted(sense))})" if sense else ""
        print(f"  {nom:8} conservada: {copiades} cel·les copiades{avis}")
    for nom in wb_v.sheetnames:
        if nom not in wb_n.sheetnames and nom != HOJA_INSTRUCCIONS:
            print(f"  {nom:8} DESAPAREIX (ja no és al JSON)")

    pds = []
    if args.pd:
        for codi, modul in moduls.items():
            nom = get_hoja_label(modul["nombre"])
            trobades = glob.glob(os.path.join(pdir, f"PD_{cicle}_{codi}_*.md"))
            if nom in canvien and not trobades:
                print(f"  PD {codi}: no hi ha fitxer")
            anterior = nom_vell(wb_v[nom]) if nom in wb_v.sheetnames else None
            for p in trobades:
                text = open(p, encoding="utf-8").read()
                canvis = []
                if nom in canvien:
                    text = actualitza_pd(text, modul, anterior)
                    if text is None:
                        errors.append(os.path.basename(p))
                        print(f"  PD {os.path.basename(p)}: ERROR, no trobe la taula de RA")
                        continue
                    canvis.append("taula de RA actualitzada")
                hores = HORES_RE.search(text)
                if hores and hores.group(2) != str(modul["horas"]):
                    text = HORES_RE.sub(lambda m: f"{m.group(1)}{modul['horas']}{m.group(3)}", text, count=1)
                    canvis.append(f"hores {hores.group(2)} → {modul['horas']}")
                seccio = UC_SECCIO_RE.search(text)
                if not seccio:
                    print(f"  PD {os.path.basename(p)}: AVÍS, no trobe la secció d'estàndards de competència")
                else:
                    abans = UC_CODI_RE.findall(seccio.group(0))
                    ara = list(modul.get("UnidadesCompetenciaAcreditadas") or {})
                    if abans != ara:
                        text = UC_SECCIO_RE.sub(lambda m: seccio_uc(modul), text, count=1)
                        canvis.append(f"UC: {', '.join(abans) or '—'} → {', '.join(ara) or '—'}")
                if not dualitza(modul) and FE_SECCIO_RE.search(text):
                    text = FE_SECCIO_RE.sub("", text)
                    canvis.append("secció de FE eliminada (ja no dualitza)")
                estat = PD_ESTAT_RE.search(p)
                desti = os.path.join(pdir, f"PD_{cicle}_{codi}_{nom_fitxer_modul(modul['nombre'])}"
                                           f"_{estat.group(1) if estat else 'BORRADOR'}.md")
                if desti != p:
                    canvis.append(f"reanomenada a {os.path.basename(desti)}")
                if canvis:
                    pds.append((p, text, desti))
                    print(f"  PD {os.path.basename(p)}: {'; '.join(canvis)}")

    if errors:
        sys.exit(f"\nNo s'ha escrit res: {len(errors)} error(s).")
    if not args.aplica:
        print("\nSimulació: no s'ha escrit res. Torna-ho a executar amb --aplica.")
        return

    marca = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    bdir = os.path.join(args.backup_dir, f"{cicle}_{marca}")
    os.makedirs(bdir, exist_ok=True)
    shutil.copy2(llibre, bdir)
    for p, _, _ in pds:
        shutil.copy2(p, bdir)
    wb_n.save(nou_path)
    shutil.copyfile(nou_path, llibre)
    for p, text, desti in pds:
        with open(p, "w", encoding="utf-8") as f:
            f.write(text)
        if desti != p:
            os.replace(p, desti)
    shutil.rmtree(tmp, ignore_errors=True)
    print(f"\nFet. Còpia dels originals a {bdir}")


if __name__ == "__main__":
    main()
